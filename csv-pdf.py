#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Convertidor recursivo:
  - Excel/Calc -> CSV
  - Word/Writer -> PDF

Uso normal:
    python convertir_office_a_csv_pdf.py

Uso no interactivo recomendado:
    python convertir_office_a_csv_pdf.py . --convert both --output both --non-interactive

Dependencias Python recomendadas para Excel:
    pip install pandas openpyxl xlrd pyxlsb odfpy python-calamine

Muy recomendado para máxima compatibilidad:
    Instalar LibreOffice. El script detecta automaticamente 'soffice' o 'libreoffice'.

Notas honestas:
  - CSV no puede conservar colores, filtros, formulas vivas, graficos, tablas dinamicas,
    comentarios, varias hojas en un solo archivo ni formato visual. Por eso cada hoja
    se exporta a un CSV independiente.
  - PDF conserva mejor la apariencia del documento, pero depende del motor de conversion
    instalado, normalmente LibreOffice o Microsoft Word en Windows.
  - Archivos cifrados con contraseña, muy corruptos, accesos directos .gdoc/.gsheet,
    archivos incompletos o formatos no soportados por el motor instalado no se pueden
    convertir magicamente. El script no se detiene: registra el error y, si hay carpeta
    central, copia el original problemático en NO_CONVERTIDOS para que no se pierda.
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import hashlib
import json
import math
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import warnings
import zipfile
import io
import textwrap
import xml.etree.ElementTree as ET
from html import unescape
from html.parser import HTMLParser
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Literal, Sequence


warnings.filterwarnings(
    "ignore",
    message=r"Workbook contains no default style.*",
    category=UserWarning,
)


# =========================
# Configuracion de formatos
# =========================

EXCEL_EXTENSIONS = {
    ".xlsx", ".xlsm", ".xltx", ".xltm",
    ".xls", ".xlt",
    ".xlsb",
    ".ods", ".ots", ".fods",
    ".xml", ".slk", ".dif", ".dbf",
}

OOXML_EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
OLD_XLS_EXTENSIONS = {".xls", ".xlt"}
XLSB_EXTENSIONS = {".xlsb"}
ODF_SPREADSHEET_EXTENSIONS = {".ods", ".ots", ".fods"}

# Documentos tipo Word/Writer. Incluye varios formatos que LibreOffice puede abrir.
# .docs se incluye porque mucha gente lo escribe asi, aunque no es un formato Word estandar.
DOCUMENT_EXTENSIONS = {
    ".doc", ".docs", ".docx", ".docm",
    ".dot", ".dotx", ".dotm",
    ".odt", ".ott", ".fodt", ".sxw",
    ".rtf",
    ".wps", ".wpt", ".wpd", ".wri",
    ".abw", ".zabw",
    ".hwp", ".lwp",
    ".txt", ".text",
    ".html", ".htm", ".xhtml",
    ".xml",
    ".pages",       # intento; suele requerir exportar desde Apple Pages
    ".gdoc",        # atajo de Google Docs: no contiene el documento real
}

GOOGLE_SHORTCUT_EXTENSIONS = {".gdoc", ".gsheet", ".gslides"}
TEMP_PREFIXES = ("~$", ".~lock.")

DEFAULT_ENCODING = "utf-8-sig"
DEFAULT_DELIMITER = ","
DEFAULT_TIMEOUT_SECONDS = 240
CENTRAL_FOLDER_PREFIX = "CONVERSIONES_OFFICE"
REPORT_BASENAME = "reporte_conversion"
ERROR_LOG_BASENAME = "errores_conversion"

JobKind = Literal["excel", "docs"]
ConvertChoice = Literal["excel", "docs", "both"]
OutputMode = Literal["same", "central", "both"]
CentralLayout = Literal["flat", "tree"]


# =========================
# Modelos de datos
# =========================

@dataclass
class RunConfig:
    root_dir: Path
    convert_choice: ConvertChoice
    output_mode: OutputMode
    delimiter: str = DEFAULT_DELIMITER
    encoding: str = DEFAULT_ENCODING
    overwrite: bool = False
    safe_csv_mode: bool = False
    prefer_libreoffice_excel: bool = False
    disable_libreoffice: bool = False
    prefer_ms_word: bool = False
    timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS
    central_layout: CentralLayout = "flat"
    pdfa: Literal["none", "1", "2", "3"] = "none"
    copy_failed_originals: bool = True
    non_interactive: bool = False


@dataclass
class SheetData:
    source_file: Path
    sheet_name: str
    rows: Iterable[Sequence[Any]]
    method: str
    warnings: list[str] = field(default_factory=list)


@dataclass
class JobFile:
    kind: JobKind
    path: Path


@dataclass
class OutputLocations:
    central_root: Path | None = None
    csv_dir: Path | None = None
    pdf_dir: Path | None = None
    reports_dir: Path | None = None
    failed_dir: Path | None = None


@dataclass
class ConversionRecord:
    kind: str
    source_file: str
    source_ext: str
    item_name: str = ""
    status: str = ""
    method: str = ""
    rows: int = 0
    columns: int = 0
    output_same_location: str = ""
    output_central_location: str = ""
    failed_original_copy: str = ""
    warnings: str = ""
    error: str = ""


class ConversionError(Exception):
    """Error controlado de conversion."""


# =========================
# Utilidades generales
# =========================

def now_stamp() -> str:
    return _dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def print_header(title: str) -> None:
    print("\n" + "=" * 72)
    print(title)
    print("=" * 72)


def sanitize_filename(value: str, fallback: str = "archivo", max_len: int = 140) -> str:
    value = str(value).strip() or fallback
    value = re.sub(r"[\\/:*?\"<>|]+", "_", value)
    value = re.sub(r"[\x00-\x1f]+", "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    if not value:
        value = fallback

    reserved = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
    if value.upper() in reserved:
        value = f"_{value}"

    if len(value) > max_len:
        digest = hashlib.sha1(value.encode("utf-8", errors="ignore")).hexdigest()[:10]
        value = f"{value[:max_len - 12]}__{digest}"
    return value


def unique_path(path: Path, overwrite: bool = False) -> Path:
    if overwrite or not path.exists():
        return path

    stem = path.stem
    suffix = path.suffix
    parent = path.parent
    for i in range(2, 100000):
        candidate = parent / f"{stem}_{i}{suffix}"
        if not candidate.exists():
            return candidate
    raise ConversionError(f"No se pudo crear un nombre unico para: {path}")


def copy_file_unique(src: Path, dst: Path, overwrite: bool = False) -> Path:
    dst.parent.mkdir(parents=True, exist_ok=True)
    final_dst = unique_path(dst, overwrite=overwrite)
    shutil.copy2(src, final_dst)
    return final_dst


def file_uri(path: Path) -> str:
    return path.resolve().as_uri()


def relative_to_or_name(path: Path, root: Path) -> Path:
    try:
        return path.resolve().relative_to(root.resolve())
    except Exception:
        return Path(path.name)


def clean_inner_tabular_extension(stem: str) -> str:
    """Evita salidas tipo archivo.csv.csv cuando el original era archivo.csv.xls."""
    cleaned = stem
    for inner_ext in (".csv", ".tsv", ".txt", ".html", ".htm", ".xml"):
        if cleaned.lower().endswith(inner_ext):
            cleaned = cleaned[: -len(inner_ext)]
            break
    return cleaned or stem


def build_flat_central_stem(root_dir: Path, source: Path, max_len: int = 180) -> str:
    rel = relative_to_or_name(source, root_dir).with_suffix("")
    if source.suffix.lower() in EXCEL_EXTENSIONS:
        rel = rel.with_name(clean_inner_tabular_extension(rel.name))
    parts = [sanitize_filename(part, max_len=80) for part in rel.parts]
    joined = "__".join(parts) or sanitize_filename(clean_inner_tabular_extension(source.stem))

    if len(joined) > max_len:
        digest = hashlib.sha1(str(rel).encode("utf-8", errors="ignore")).hexdigest()[:12]
        joined = f"{joined[:max_len - 14]}__{digest}"
    return joined


def build_tree_central_path(base_dir: Path, root_dir: Path, source: Path, suffix: str, extra_stem: str = "") -> Path:
    rel = relative_to_or_name(source, root_dir)
    rel_parent = Path(*[sanitize_filename(p, max_len=80) for p in rel.parent.parts]) if rel.parent.parts else Path()
    raw_stem = clean_inner_tabular_extension(source.stem) if source.suffix.lower() in EXCEL_EXTENSIONS else source.stem
    stem = sanitize_filename(raw_stem)
    if extra_stem:
        stem = f"{stem}__{sanitize_filename(extra_stem)}"
    return base_dir / rel_parent / f"{stem}{suffix}"


def build_flat_central_path(base_dir: Path, root_dir: Path, source: Path, suffix: str, extra_stem: str = "") -> Path:
    stem = build_flat_central_stem(root_dir, source)
    if extra_stem:
        stem = f"{stem}__{sanitize_filename(extra_stem)}"
    return base_dir / f"{stem}{suffix}"


def build_central_path(
    base_dir: Path,
    root_dir: Path,
    source: Path,
    suffix: str,
    extra_stem: str = "",
    layout: CentralLayout = "flat",
) -> Path:
    if layout == "tree":
        return build_tree_central_path(base_dir, root_dir, source, suffix, extra_stem=extra_stem)
    return build_flat_central_path(base_dir, root_dir, source, suffix, extra_stem=extra_stem)


def is_inside(path: Path, maybe_parent: Path | None) -> bool:
    if maybe_parent is None:
        return False
    try:
        path.resolve().relative_to(maybe_parent.resolve())
        return True
    except Exception:
        return False


def is_inside_generated_conversion_folder(path: Path) -> bool:
    """Evita reprocesar carpetas centrales de ejecuciones anteriores."""
    for parent in path.resolve().parents:
        if parent.name.startswith(f"{CENTRAL_FOLDER_PREFIX}_"):
            return True
    return False


def should_skip_path(path: Path, central_root: Path | None) -> bool:
    name = path.name
    if any(name.startswith(prefix) for prefix in TEMP_PREFIXES):
        return True
    if is_inside(path, central_root):
        return True
    if is_inside_generated_conversion_folder(path):
        return True
    return False



# =========================
# Deteccion por contenido y formatos tabulares disfrazados
# =========================

OLE_CFB_SIGNATURE = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
TEXT_CONTROL_BYTES = set(range(0, 9)) | set(range(14, 32))


def read_prefix(path: Path, size: int = 65536) -> bytes:
    try:
        with path.open("rb") as f:
            return f.read(size)
    except Exception:
        return b""


def has_known_binary_spreadsheet_signature(data: bytes) -> bool:
    return data.startswith(OLE_CFB_SIGNATURE) or data.startswith(ZIP_SIGNATURES)


def detect_text_encoding(data: bytes) -> str:
    """Detecta una codificacion razonable sin dependencias externas."""
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if data.startswith(b"\xff\xfe"):
        return "utf-16-le"
    if data.startswith(b"\xfe\xff"):
        return "utf-16-be"

    # Si hay muchos NUL alternos, probablemente es UTF-16 sin BOM.
    sample = data[:4096]
    if sample:
        even_nuls = sum(1 for i in range(0, len(sample), 2) if sample[i] == 0)
        odd_nuls = sum(1 for i in range(1, len(sample), 2) if sample[i] == 0)
        pairs = max(1, len(sample) // 2)
        if odd_nuls / pairs > 0.25 and even_nuls / pairs < 0.05:
            return "utf-16-le"
        if even_nuls / pairs > 0.25 and odd_nuls / pairs < 0.05:
            return "utf-16-be"

    candidates = ["utf-8", "cp1252", "latin-1"]
    best = "latin-1"
    best_score = float("inf")
    for enc in candidates:
        try:
            text = data.decode(enc, errors="replace")
        except Exception:
            continue
        replacement_count = text.count("\ufffd")
        control_count = sum(1 for ch in text[:4096] if ord(ch) < 32 and ch not in "\r\n\t")
        score = replacement_count * 10 + control_count
        if score < best_score:
            best = enc
            best_score = score
    return best


def read_text_lossy(path: Path, max_bytes: int | None = None) -> tuple[str, str]:
    data = path.read_bytes() if max_bytes is None else read_prefix(path, max_bytes)
    enc = detect_text_encoding(data)
    try:
        return data.decode(enc, errors="replace"), enc
    except Exception:
        return data.decode("latin-1", errors="replace"), "latin-1"


def is_probably_text_bytes(data: bytes) -> bool:
    if not data:
        return True
    if has_known_binary_spreadsheet_signature(data):
        return False
    enc = detect_text_encoding(data)
    if enc.startswith("utf-16"):
        return True
    sample = data[:8192]
    if not sample:
        return True
    control = sum(1 for b in sample if b in TEXT_CONTROL_BYTES)
    return (control / max(1, len(sample))) < 0.08


def xml_local_name(tag: Any) -> str:
    text = str(tag)
    if "}" in text:
        text = text.rsplit("}", 1)[-1]
    return text


def xml_attr_by_local_name(element: ET.Element, local_name: str, default: Any = None) -> Any:
    for key, value in element.attrib.items():
        if xml_local_name(key).lower() == local_name.lower():
            return value
    return default


def xml_children(element: ET.Element, local_name: str) -> list[ET.Element]:
    lname = local_name.lower()
    return [child for child in list(element) if xml_local_name(child.tag).lower() == lname]


def xml_first_descendant(element: ET.Element, local_name: str) -> ET.Element | None:
    lname = local_name.lower()
    for child in element.iter():
        if xml_local_name(child.tag).lower() == lname:
            return child
    return None


def normalize_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def detect_delimiter_from_text(text: str) -> str | None:
    text = normalize_newlines(text)
    lines = text.split("\n")
    non_empty = [line for line in lines[:80] if line.strip()]
    if not non_empty:
        return None

    first = non_empty[0].strip()
    if len(first) >= 5 and first.lower().startswith("sep="):
        return first[4:5]

    sample = "\n".join(non_empty[:30])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        if dialect.delimiter:
            return dialect.delimiter
    except Exception:
        pass

    candidates = [",", ";", "\t", "|"]
    scores: list[tuple[float, str]] = []
    for delimiter in candidates:
        counts = [line.count(delimiter) for line in non_empty[:30]]
        positive = [c for c in counts if c > 0]
        if not positive:
            continue
        # Premia delimitadores repetidos y razonablemente consistentes.
        avg = sum(positive) / len(positive)
        consistency = len(positive) / max(1, len(counts))
        variance_penalty = (max(positive) - min(positive)) * 0.05 if len(positive) > 1 else 0
        scores.append((avg * consistency - variance_penalty, delimiter))

    if not scores:
        return None
    scores.sort(reverse=True)
    if scores[0][0] <= 0:
        return None
    return scores[0][1]


def sniff_text_tabular_kind(path: Path) -> str | None:
    """Devuelve html, xml_spreadsheet, delimited o None segun el contenido real."""
    data = read_prefix(path, 131072)
    if not data or not is_probably_text_bytes(data):
        return None
    text, _ = read_text_lossy(path, max_bytes=131072)
    sample = normalize_newlines(text).lstrip("\ufeff\x00\t \r\n")
    low = sample[:20000].lower()

    if "<table" in low or "<html" in low or "<!doctype html" in low:
        return "html"
    if "<workbook" in low and "spreadsheet" in low:
        return "xml_spreadsheet"
    if "urn:schemas-microsoft-com:office:spreadsheet" in low:
        return "xml_spreadsheet"
    if detect_delimiter_from_text(sample) is not None:
        return "delimited"
    # Muchos reportes vienen como una sola columna HTML/XML mal rotulado. No se fuerza aqui
    # para evitar convertir basura binaria o XML no tabular.
    return None


class SimpleHTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table_depth = 0
        self._current_table: list[list[str]] | None = None
        self._current_row: list[str] | None = None
        self._cell_chunks: list[str] | None = None
        self._cell_colspan = 1

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag_l = tag.lower()
        attrs_dict = {k.lower(): v for k, v in attrs}
        if tag_l == "table":
            if self._table_depth == 0:
                self._current_table = []
            self._table_depth += 1
        elif self._table_depth > 0 and tag_l == "tr":
            self._current_row = []
        elif self._table_depth > 0 and self._current_row is not None and tag_l in {"td", "th"}:
            self._cell_chunks = []
            try:
                self._cell_colspan = max(1, int(attrs_dict.get("colspan") or 1))
            except Exception:
                self._cell_colspan = 1
        elif self._cell_chunks is not None and tag_l in {"br", "p", "div"}:
            self._cell_chunks.append(" ")

    def handle_data(self, data: str) -> None:
        if self._cell_chunks is not None:
            self._cell_chunks.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag_l = tag.lower()
        if tag_l in {"td", "th"} and self._cell_chunks is not None and self._current_row is not None:
            text = re.sub(r"\s+", " ", unescape("".join(self._cell_chunks))).strip()
            for _ in range(self._cell_colspan):
                self._current_row.append(text)
            self._cell_chunks = None
            self._cell_colspan = 1
        elif tag_l == "tr" and self._current_row is not None:
            if self._current_table is not None and any(cell != "" for cell in self._current_row):
                self._current_table.append(self._current_row)
            self._current_row = None
        elif tag_l == "table" and self._table_depth > 0:
            self._table_depth -= 1
            if self._table_depth == 0 and self._current_table is not None:
                if self._current_table:
                    self.tables.append(self._current_table)
                self._current_table = None


def iter_html_table_sheets(source: Path) -> Iterator[SheetData]:
    text, enc = read_text_lossy(source)
    low = text[:20000].lower()
    if "<table" not in low and "<html" not in low and "<!doctype html" not in low:
        raise ConversionError("El archivo no parece HTML tabular.")

    parser = SimpleHTMLTableParser()
    try:
        parser.feed(text)
    except Exception:
        # HTMLParser es tolerante, pero si se rompe continuamos con pandas si existe.
        pass

    if parser.tables:
        for idx, rows in enumerate(parser.tables, start=1):
            yield SheetData(
                source_file=source,
                sheet_name=f"Tabla{idx}",
                rows=rows,
                method=f"html-table-parser:{enc}",
                warnings=["Archivo HTML guardado con extension de Excel; se extrajeron tablas HTML."],
            )
        return

    try:
        import pandas as pd
        tables = pd.read_html(io.StringIO(text), keep_default_na=False)
    except Exception as exc:
        raise ConversionError(f"No se pudieron extraer tablas HTML: {exc}") from exc

    if not tables:
        raise ConversionError("El HTML no contiene tablas convertibles.")

    for idx, df in enumerate(tables, start=1):
        try:
            df = df.where(df.notna(), "")
        except Exception:
            pass
        rows: list[list[Any]] = []
        # Si pandas interpreto <th> como columnas, reincorporamos esa fila para no perder encabezados.
        try:
            import pandas as pd  # type: ignore
            is_default_columns = isinstance(df.columns, pd.RangeIndex)
        except Exception:
            is_default_columns = False
        if not is_default_columns:
            rows.append([str(col) for col in list(df.columns)])
        rows.extend([list(row) for row in df.itertuples(index=False, name=None)])
        yield SheetData(
            source_file=source,
            sheet_name=f"Tabla{idx}",
            rows=rows,
            method=f"pandas.read_html:{enc}",
            warnings=["Archivo HTML guardado con extension de Excel; se extrajeron tablas HTML."],
        )


def iter_xml_spreadsheet_sheets(source: Path) -> Iterator[SheetData]:
    text, enc = read_text_lossy(source)
    low = text[:50000].lower()
    if "spreadsheet" not in low and "<workbook" not in low:
        raise ConversionError("El XML no parece SpreadsheetML de Excel 2003.")

    try:
        # Usar bytes permite que ElementTree respete la declaracion de encoding del XML.
        root = ET.fromstring(source.read_bytes())
    except Exception:
        try:
            # Fallback para XML servido con encoding erroneo.
            cleaned_text = re.sub(r"<\?xml[^>]*encoding=[\"'][^\"']+[\"'][^>]*\?>", "", text, count=1, flags=re.I).lstrip()
            root = ET.fromstring(cleaned_text)
        except Exception as exc:
            raise ConversionError(f"No se pudo leer XML Spreadsheet: {exc}") from exc

    worksheets = [el for el in root.iter() if xml_local_name(el.tag).lower() == "worksheet"]
    if not worksheets:
        raise ConversionError("El XML Spreadsheet no contiene Worksheet.")

    for ws_idx, worksheet in enumerate(worksheets, start=1):
        sheet_name = xml_attr_by_local_name(worksheet, "Name", f"Hoja{ws_idx}") or f"Hoja{ws_idx}"
        table = None
        for child in worksheet.iter():
            if xml_local_name(child.tag).lower() == "table":
                table = child
                break
        if table is None:
            yield SheetData(source, str(sheet_name), [], method=f"xml-spreadsheet:{enc}", warnings=["Worksheet sin tabla."])
            continue

        rows: list[list[Any]] = []
        current_row_number = 1
        for row_el in xml_children(table, "Row"):
            idx_raw = xml_attr_by_local_name(row_el, "Index")
            try:
                row_index = int(float(idx_raw)) if idx_raw else current_row_number
            except Exception:
                row_index = current_row_number
            while current_row_number < row_index:
                rows.append([])
                current_row_number += 1

            row_values: list[Any] = []
            current_col_number = 1
            for cell_el in xml_children(row_el, "Cell"):
                cidx_raw = xml_attr_by_local_name(cell_el, "Index")
                try:
                    col_index = int(float(cidx_raw)) if cidx_raw else current_col_number
                except Exception:
                    col_index = current_col_number
                while current_col_number < col_index:
                    row_values.append("")
                    current_col_number += 1

                data_el = xml_first_descendant(cell_el, "Data")
                value = "" if data_el is None else "".join(data_el.itertext())
                row_values.append(value)
                current_col_number += 1

            rows.append(row_values)
            current_row_number += 1

        yield SheetData(
            source_file=source,
            sheet_name=str(sheet_name),
            rows=rows,
            method=f"xml-spreadsheet:{enc}",
            warnings=["Archivo XML SpreadsheetML guardado con extension de Excel."],
        )


def iter_text_delimited_sheets(source: Path) -> Iterator[SheetData]:
    data = read_prefix(source, 131072)
    if not is_probably_text_bytes(data):
        raise ConversionError("El archivo no parece texto delimitado.")

    text, enc = read_text_lossy(source)
    text = normalize_newlines(text).lstrip("\ufeff")
    delimiter = detect_delimiter_from_text(text)
    if delimiter is None:
        raise ConversionError("No se detecto delimitador CSV/TSV confiable.")

    lines = text.split("\n")
    if lines and lines[0].strip().lower().startswith("sep="):
        text = "\n".join(lines[1:])

    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [row for row in reader]
    except Exception as exc:
        raise ConversionError(f"No se pudo leer como texto delimitado: {exc}") from exc

    # Evita falsos positivos de texto de una sola celda y sin separadores reales.
    non_empty_rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if not non_empty_rows:
        raise ConversionError("El texto delimitado esta vacio.")
    if max((len(row) for row in non_empty_rows), default=0) <= 1 and len(non_empty_rows) <= 1:
        raise ConversionError("El archivo no parece una tabla CSV/TSV suficiente.")

    yield SheetData(
        source_file=source,
        sheet_name="CSV_TSV",
        rows=rows,
        method=f"text-delimited:{enc}:delimiter={repr(delimiter)}",
        warnings=["Archivo de texto delimitado guardado con extension de Excel."],
    )


# =========================
# Deteccion de motores
# =========================

def get_soffice_executable() -> str | None:
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found

    common_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/snap/bin/libreoffice",
        "/opt/homebrew/bin/soffice",
    ]
    for item in common_paths:
        if Path(item).exists():
            return item
    return None


def microsoft_word_com_available() -> bool:
    if platform.system().lower() != "windows":
        return False
    try:
        import win32com.client  # type: ignore  # noqa: F401
        return True
    except Exception:
        return False


# =========================
# LibreOffice comun
# =========================

def build_pdf_filter(pdfa: Literal["none", "1", "2", "3"] = "none") -> str:
    # writer_pdf_Export funciona para documentos de Writer. Para otros modulos, LO suele resolverlo,
    # pero aqui convertimos documentos tipo Word/Writer.
    if pdfa == "none":
        return "pdf:writer_pdf_Export"
    version_map = {"1": 1, "2": 2, "3": 3}
    options = {"SelectPdfVersion": {"type": "long", "value": str(version_map[pdfa])}}
    return "pdf:writer_pdf_Export:" + json.dumps(options, separators=(",", ":"))


def run_libreoffice_convert(
    source: Path,
    output_dir: Path,
    convert_to: str,
    expected_suffix: str,
    soffice: str,
    timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    profile_dir = output_dir / "_lo_profile"
    profile_dir.mkdir(parents=True, exist_ok=True)

    before = {p.resolve() for p in output_dir.glob(f"*{expected_suffix}")}

    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--nodefault",
        "--nofirststartwizard",
        "--nolockcheck",
        f"-env:UserInstallation={file_uri(profile_dir)}",
        "--convert-to", convert_to,
        "--outdir", str(output_dir),
        str(source),
    ]

    try:
        completed = subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=timeout_seconds,
        )
    except subprocess.TimeoutExpired as exc:
        raise ConversionError(f"LibreOffice excedio el tiempo limite de {timeout_seconds}s") from exc
    except Exception as exc:
        raise ConversionError(f"No se pudo ejecutar LibreOffice: {exc}") from exc

    expected = output_dir / f"{source.stem}{expected_suffix}"
    candidates = [p for p in output_dir.glob(f"*{expected_suffix}") if p.resolve() not in before and p.is_file()]
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)

    if expected.exists() and expected.stat().st_size > 0:
        return expected
    if candidates:
        for candidate in candidates:
            if candidate.stat().st_size > 0:
                return candidate

    # A veces LibreOffice sobrescribe un archivo que ya existia en el outdir.
    all_candidates = sorted(output_dir.glob(f"*{expected_suffix}"), key=lambda p: p.stat().st_mtime, reverse=True)
    for candidate in all_candidates:
        if candidate.is_file() and candidate.stat().st_size > 0:
            return candidate

    stdout = (completed.stdout or "").strip()
    stderr = (completed.stderr or "").strip()
    code = completed.returncode
    raise ConversionError(
        "LibreOffice no genero salida. "
        f"Codigo={code}. STDOUT={stdout[:800]!r}. STDERR={stderr[:800]!r}"
    )


def validate_pdf(path: Path) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise ConversionError("El PDF generado esta vacio o no existe.")
    try:
        with path.open("rb") as f:
            head = f.read(5)
        if head != b"%PDF-":
            raise ConversionError("El archivo generado no parece ser un PDF valido.")
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"No se pudo validar el PDF generado: {exc}") from exc


def libreoffice_convert_to_xlsx(source: Path, temp_root: Path, soffice: str, timeout_seconds: int) -> Path:
    outdir = temp_root / "lo_xlsx"
    try:
        result = run_libreoffice_convert(
            source=source,
            output_dir=outdir,
            convert_to="xlsx",
            expected_suffix=".xlsx",
            soffice=soffice,
            timeout_seconds=timeout_seconds,
        )
    except Exception as exc:
        # Segundo intento con nombre ASCII corto por problemas de rutas/caracteres.
        safe_input_dir = temp_root / "lo_xlsx_safe_input"
        safe_input_dir.mkdir(parents=True, exist_ok=True)
        safe_source = safe_input_dir / f"input{source.suffix.lower()}"
        shutil.copy2(source, safe_source)
        try:
            result = run_libreoffice_convert(
                source=safe_source,
                output_dir=outdir,
                convert_to="xlsx",
                expected_suffix=".xlsx",
                soffice=soffice,
                timeout_seconds=timeout_seconds,
            )
        except Exception as exc2:
            raise ConversionError(f"LibreOffice no pudo normalizar a XLSX. Intento 1: {exc}; intento 2: {exc2}") from exc2
    return result


def libreoffice_convert_document_to_pdf(
    source: Path,
    temp_root: Path,
    soffice: str,
    timeout_seconds: int,
    pdfa: Literal["none", "1", "2", "3"] = "none",
) -> Path:
    outdir = temp_root / "lo_pdf"
    pdf_filter = build_pdf_filter(pdfa)

    errors: list[str] = []
    for attempt, input_path in enumerate([source, None], start=1):
        if input_path is None:
            safe_input_dir = temp_root / "lo_pdf_safe_input"
            safe_input_dir.mkdir(parents=True, exist_ok=True)
            input_path = safe_input_dir / f"input{source.suffix.lower()}"
            shutil.copy2(source, input_path)

        try:
            result = run_libreoffice_convert(
                source=input_path,
                output_dir=outdir,
                convert_to=pdf_filter,
                expected_suffix=".pdf",
                soffice=soffice,
                timeout_seconds=timeout_seconds,
            )
            validate_pdf(result)
            return result
        except Exception as exc:
            errors.append(f"intento {attempt}: {exc}")

    raise ConversionError("LibreOffice no pudo convertir a PDF. " + " | ".join(errors))


# =========================
# Microsoft Word COM opcional
# =========================

def ms_word_convert_to_pdf(
    source: Path,
    temp_root: Path,
    pdfa: Literal["none", "1", "2", "3"] = "none",
) -> Path:
    if platform.system().lower() != "windows":
        raise ConversionError("Microsoft Word COM solo esta disponible en Windows.")

    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except Exception as exc:
        raise ConversionError("Falta pywin32 para usar Microsoft Word COM: pip install pywin32") from exc

    outdir = temp_root / "ms_word_pdf"
    outdir.mkdir(parents=True, exist_ok=True)
    out_pdf = outdir / f"{sanitize_filename(source.stem)}.pdf"

    word = None
    doc = None
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        doc = word.Documents.Open(
            str(source.resolve()),
            ReadOnly=True,
            AddToRecentFiles=False,
            ConfirmConversions=False,
            NoEncodingDialog=True,
        )

        # ExportFormat=17 => wdExportFormatPDF.
        # UseISO19005_1 solo produce PDF/A-1; para PDF/A-2/3 se recomienda LibreOffice.
        use_pdfa_1 = pdfa == "1"
        doc.ExportAsFixedFormat(
            OutputFileName=str(out_pdf.resolve()),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=use_pdfa_1,
        )
        validate_pdf(out_pdf)
        return out_pdf

    except Exception as exc:
        raise ConversionError(f"Microsoft Word COM no pudo convertir a PDF: {exc}") from exc
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()  # type: ignore[name-defined]
        except Exception:
            pass


# =========================
# Excel -> CSV
# =========================

def normalize_cell_value(value: Any, safe_csv_mode: bool = False) -> str:
    if value is None:
        text = ""
    elif isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, _dt.datetime):
        text = value.isoformat(sep=" ", timespec="seconds")
    elif isinstance(value, _dt.date):
        text = value.isoformat()
    elif isinstance(value, _dt.time):
        text = value.isoformat(timespec="seconds")
    elif isinstance(value, Decimal):
        text = format(value, "f")
    elif isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            text = str(value)
        else:
            text = format(value, ".15g")
    else:
        text = str(value)

    # Proteccion opcional contra CSV/Formula Injection. No se activa por defecto porque cambia texto.
    if safe_csv_mode and text and text[0] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        text = "'" + text
    return text


def write_rows_to_csv(
    rows: Iterable[Sequence[Any]],
    output_path: Path,
    delimiter: str,
    encoding: str,
    safe_csv_mode: bool,
) -> tuple[int, int]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    row_count = 0
    max_cols = 0

    with output_path.open("w", newline="", encoding=encoding) as f:
        writer = csv.writer(
            f,
            delimiter=delimiter,
            quotechar='"',
            quoting=csv.QUOTE_MINIMAL,
            lineterminator="\n",
        )
        for row in rows:
            normalized = [normalize_cell_value(v, safe_csv_mode=safe_csv_mode) for v in row]
            writer.writerow(normalized)
            row_count += 1
            max_cols = max(max_cols, len(normalized))

    return row_count, max_cols


def iter_openpyxl_sheets(source: Path, method: str = "openpyxl") -> Iterator[SheetData]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConversionError("Falta instalar openpyxl: pip install openpyxl") from exc

    try:
        wb_values = load_workbook(source, data_only=True, read_only=False)
        wb_formula = load_workbook(source, data_only=False, read_only=False)
    except Exception as exc:
        raise ConversionError(f"openpyxl no pudo abrir el archivo: {exc}") from exc

    try:
        for ws_values in wb_values.worksheets:
            ws_formula = wb_formula[ws_values.title]
            warnings: list[str] = []

            merged_map: dict[tuple[int, int], Any] = {}
            merged_total = 0
            try:
                merged_ranges = list(ws_values.merged_cells.ranges)
            except Exception:
                merged_ranges = []

            for merged in merged_ranges:
                top_value = ws_values.cell(merged.min_row, merged.min_col).value
                area = (merged.max_row - merged.min_row + 1) * (merged.max_col - merged.min_col + 1)
                merged_total += area
                if merged_total <= 1_000_000:
                    for r in range(merged.min_row, merged.max_row + 1):
                        for c in range(merged.min_col, merged.max_col + 1):
                            merged_map[(r, c)] = top_value

            if merged_total > 1_000_000:
                warnings.append("Hay demasiadas celdas combinadas; se expandio solo una parte para evitar uso excesivo de memoria.")
            elif merged_total:
                warnings.append("Las celdas combinadas se expandieron repitiendo el valor superior izquierdo.")

            max_row = 0
            max_col = 0
            formula_without_cache = 0
            formula_with_cache = 0

            # Primer barrido: limites reales y formulas.
            for row in ws_values.iter_rows():
                for cell in row:
                    value = merged_map.get((cell.row, cell.column), cell.value)
                    try:
                        formula_value = ws_formula.cell(cell.row, cell.column).value
                    except Exception:
                        formula_value = None

                    if isinstance(formula_value, str) and formula_value.startswith("="):
                        if value is None:
                            formula_without_cache += 1
                            value = formula_value
                        else:
                            formula_with_cache += 1

                    if value not in (None, ""):
                        max_row = max(max_row, cell.row)
                        max_col = max(max_col, cell.column)

            if formula_with_cache:
                warnings.append(f"Se exportaron valores cacheados de {formula_with_cache} formula/s.")
            if formula_without_cache:
                warnings.append(f"{formula_without_cache} formula/s no tenian valor cacheado; se exporto el texto de la formula.")

            rows_snapshot: list[list[Any]] = []
            if max_row and max_col:
                for r in range(1, max_row + 1):
                    out: list[Any] = []
                    for c in range(1, max_col + 1):
                        value = merged_map.get((r, c), ws_values.cell(r, c).value)
                        try:
                            formula_value = ws_formula.cell(r, c).value
                        except Exception:
                            formula_value = None
                        if isinstance(formula_value, str) and formula_value.startswith("=") and value is None:
                            value = formula_value
                        out.append(value)
                    rows_snapshot.append(out)

            yield SheetData(source_file=source, sheet_name=ws_values.title, rows=rows_snapshot, method=method, warnings=warnings)

    finally:
        try:
            wb_values.close()
        except Exception:
            pass
        try:
            wb_formula.close()
        except Exception:
            pass


def iter_xlrd_sheets(source: Path) -> Iterator[SheetData]:
    try:
        import xlrd
    except ImportError as exc:
        raise ConversionError("Falta instalar xlrd: pip install xlrd") from exc

    try:
        try:
            book = xlrd.open_workbook(str(source), formatting_info=True, on_demand=True)
        except Exception:
            book = xlrd.open_workbook(str(source), formatting_info=False, on_demand=True)
    except Exception as exc:
        raise ConversionError(f"xlrd no pudo abrir el archivo .xls: {exc}") from exc

    def cell_to_value(sheet: Any, row_idx: int, col_idx: int) -> Any:
        cell = sheet.cell(row_idx, col_idx)
        ctype = cell.ctype
        value = cell.value
        if ctype == xlrd.XL_CELL_EMPTY or ctype == xlrd.XL_CELL_BLANK:
            return ""
        if ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate.xldate_as_datetime(value, book.datemode)
            except Exception:
                return value
        if ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(value)
        if ctype == xlrd.XL_CELL_ERROR:
            try:
                return xlrd.biffh.error_text_from_code.get(value, f"#ERROR_{value}")
            except Exception:
                return f"#ERROR_{value}"
        return value

    try:
        for sheet_name in book.sheet_names():
            sheet = book.sheet_by_name(sheet_name)
            warnings: list[str] = []
            merged_map: dict[tuple[int, int], Any] = {}

            merged_ranges = getattr(sheet, "merged_cells", []) or []
            merged_total = 0
            for rlo, rhi, clo, chi in merged_ranges:
                top_value = cell_to_value(sheet, rlo, clo)
                area = (rhi - rlo) * (chi - clo)
                merged_total += area
                if merged_total <= 1_000_000:
                    for r in range(rlo, rhi):
                        for c in range(clo, chi):
                            merged_map[(r, c)] = top_value

            if merged_total > 1_000_000:
                warnings.append("Hay demasiadas celdas combinadas; se expandio solo una parte para evitar uso excesivo de memoria.")
            elif merged_total:
                warnings.append("Las celdas combinadas se expandieron repitiendo el valor superior izquierdo.")

            max_row = 0
            max_col = 0
            for r in range(sheet.nrows):
                for c in range(sheet.ncols):
                    value = merged_map.get((r, c), cell_to_value(sheet, r, c))
                    if value not in (None, ""):
                        max_row = max(max_row, r + 1)
                        max_col = max(max_col, c + 1)

            rows_snapshot = [
                [merged_map.get((r, c), cell_to_value(sheet, r, c)) for c in range(max_col)]
                for r in range(max_row)
            ]
            yield SheetData(source_file=source, sheet_name=sheet_name, rows=rows_snapshot, method="xlrd", warnings=warnings)

    finally:
        try:
            book.release_resources()
        except Exception:
            pass


def iter_xlsb_sheets(source: Path) -> Iterator[SheetData]:
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:
        raise ConversionError("Falta instalar pyxlsb: pip install pyxlsb") from exc

    try:
        with open_workbook(str(source)) as wb:
            sheet_names = list(wb.sheets)
    except Exception as exc:
        raise ConversionError(f"pyxlsb no pudo abrir el archivo .xlsb: {exc}") from exc

    for sheet_name in sheet_names:
        try:
            max_row = 0
            max_col = 0
            rows_snapshot: list[list[Any]] = []

            with open_workbook(str(source)) as wb:
                with wb.get_sheet(sheet_name) as sh:
                    for r_idx, row in enumerate(sh.rows(), start=1):
                        values = [cell.v for cell in row]
                        last_col = 0
                        for i, value in enumerate(values, start=1):
                            if value not in (None, ""):
                                last_col = i
                        if last_col:
                            max_row = r_idx
                            max_col = max(max_col, last_col)
                        rows_snapshot.append(values)

            if max_row == 0 or max_col == 0:
                trimmed: list[list[Any]] = []
            else:
                trimmed = []
                for values in rows_snapshot[:max_row]:
                    if len(values) < max_col:
                        values = values + [""] * (max_col - len(values))
                    trimmed.append(values[:max_col])

            yield SheetData(
                source_file=source,
                sheet_name=sheet_name,
                rows=trimmed,
                method="pyxlsb",
                warnings=["Formato XLSB: se exportan valores; no se conserva formato visual."],
            )
        except Exception as exc:
            raise ConversionError(f"pyxlsb fallo leyendo la hoja {sheet_name!r}: {exc}") from exc


def iter_pandas_sheets(source: Path, engine: str | None = None) -> Iterator[SheetData]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise ConversionError("Falta instalar pandas: pip install pandas") from exc

    try:
        excel = pd.ExcelFile(source, engine=engine)
    except Exception as exc:
        engine_text = engine or "auto"
        raise ConversionError(f"pandas no pudo abrir el archivo con engine={engine_text}: {exc}") from exc

    try:
        for sheet_name in excel.sheet_names:
            try:
                df = excel.parse(
                    sheet_name=sheet_name,
                    header=None,
                    dtype=object,
                    keep_default_na=False,
                    na_filter=False,
                )
                df = df.where(df.notna(), "")
                rows = list(df.itertuples(index=False, name=None))
                yield SheetData(
                    source_file=source,
                    sheet_name=str(sheet_name),
                    rows=rows,
                    method=f"pandas:{engine or 'auto'}",
                    warnings=["Conversion de respaldo con pandas; puede no preservar celdas combinadas ni formato visual."],
                )
            except Exception as exc:
                raise ConversionError(f"pandas fallo leyendo la hoja {sheet_name!r}: {exc}") from exc
    finally:
        try:
            excel.close()
        except Exception:
            pass


def openpyxl_formula_cache_warning(source: Path) -> str | None:
    try:
        from openpyxl import load_workbook
    except Exception:
        return None

    try:
        wb_formula = load_workbook(source, data_only=False, read_only=True)
        wb_values = load_workbook(source, data_only=True, read_only=True)
    except Exception:
        return None

    missing = 0
    total_formula = 0
    try:
        for ws_formula in wb_formula.worksheets:
            try:
                ws_values = wb_values[ws_formula.title]
            except Exception:
                continue
            for row_formula, row_values in zip(ws_formula.iter_rows(), ws_values.iter_rows()):
                for cell_formula, cell_value in zip(row_formula, row_values):
                    val = cell_formula.value
                    if isinstance(val, str) and val.startswith("="):
                        total_formula += 1
                        if cell_value.value is None:
                            missing += 1
                            if missing >= 25:
                                break
                if missing >= 25:
                    break
            if missing >= 25:
                break
    finally:
        try:
            wb_formula.close()
        except Exception:
            pass
        try:
            wb_values.close()
        except Exception:
            pass

    if missing:
        return f"Se detectaron formulas sin valor cacheado ({missing} muestra/s); conviene recalcular con LibreOffice."
    if total_formula:
        return f"Se detectaron {total_formula} formula/s con valores cacheados."
    return None


def iter_excel_sheets_with_fallbacks(
    source: Path,
    soffice: str | None,
    timeout_seconds: int,
    prefer_libreoffice: bool = False,
) -> Iterator[SheetData]:
    ext = source.suffix.lower()
    errors: list[str] = []

    with tempfile.TemporaryDirectory(prefix="office_excel_to_csv_") as tmp:
        temp_root = Path(tmp)

        if prefer_libreoffice and soffice:
            try:
                normalized = libreoffice_convert_to_xlsx(source, temp_root, soffice, timeout_seconds=timeout_seconds)
                for sheet in iter_openpyxl_sheets(normalized, method="libreoffice->xlsx->openpyxl"):
                    sheet.source_file = source
                    yield sheet
                return
            except Exception as exc:
                errors.append(f"LibreOffice primero: {exc}")

        content_kind = sniff_text_tabular_kind(source)
        text_attempts: list[tuple[str, Callable[[], Iterator[SheetData]]]] = []
        if content_kind == "html":
            text_attempts.extend([
                ("html-table", lambda: iter_html_table_sheets(source)),
                ("text-delimited", lambda: iter_text_delimited_sheets(source)),
            ])
        elif content_kind == "xml_spreadsheet":
            text_attempts.extend([
                ("xml-spreadsheet", lambda: iter_xml_spreadsheet_sheets(source)),
                ("html-table", lambda: iter_html_table_sheets(source)),
                ("text-delimited", lambda: iter_text_delimited_sheets(source)),
            ])
        elif content_kind == "delimited":
            text_attempts.extend([
                ("text-delimited", lambda: iter_text_delimited_sheets(source)),
                ("html-table", lambda: iter_html_table_sheets(source)),
                ("xml-spreadsheet", lambda: iter_xml_spreadsheet_sheets(source)),
            ])

        attempted_names: set[str] = set()
        for name, attempt in text_attempts:
            attempted_names.add(name)
            try:
                sheets = list(attempt())
                if sheets:
                    for sheet in sheets:
                        yield sheet
                    return
            except Exception as exc:
                errors.append(f"{name}: {exc}")

        attempts: list[tuple[str, Callable[[], Iterator[SheetData]]]] = []
        if ext in OOXML_EXCEL_EXTENSIONS:
            attempts.append(("openpyxl", lambda: iter_openpyxl_sheets(source)))
        elif ext in OLD_XLS_EXTENSIONS:
            attempts.append(("xlrd", lambda: iter_xlrd_sheets(source)))
        elif ext in XLSB_EXTENSIONS:
            attempts.append(("pyxlsb", lambda: iter_xlsb_sheets(source)))
        elif ext in ODF_SPREADSHEET_EXTENSIONS:
            attempts.append(("pandas:odf", lambda: iter_pandas_sheets(source, engine="odf")))

        # Calamine suele abrir xls/xlsx/xlsm/xlsb/ods si pandas y python-calamine son compatibles.
        attempts.append(("pandas:calamine", lambda: iter_pandas_sheets(source, engine="calamine")))
        attempts.append(("pandas:auto", lambda: iter_pandas_sheets(source, engine=None)))

        for name, attempt in attempts:
            try:
                sheets = list(attempt())
                if sheets:
                    warning = openpyxl_formula_cache_warning(source) if ext in OOXML_EXCEL_EXTENSIONS else None

                    # Si hay formulas sin cache y tenemos LibreOffice, preferimos recalcular/normalizar.
                    if warning and "sin valor cacheado" in warning and soffice:
                        errors.append(f"{name}: {warning}")
                        raise ConversionError(warning)

                    for sheet in sheets:
                        if warning:
                            sheet.warnings.append(warning)
                        yield sheet
                    return
            except Exception as exc:
                errors.append(f"{name}: {exc}")

        # Si los motores de Excel fallaron y el archivo es texto, intenta de nuevo con
        # parsers tolerantes. Esto cubre .xls falsos generados por portales web.
        if is_probably_text_bytes(read_prefix(source, 131072)):
            late_attempts: list[tuple[str, Callable[[], Iterator[SheetData]]]] = [
                ("html-table", lambda: iter_html_table_sheets(source)),
                ("xml-spreadsheet", lambda: iter_xml_spreadsheet_sheets(source)),
                ("text-delimited", lambda: iter_text_delimited_sheets(source)),
            ]
            for name, attempt in late_attempts:
                if name in attempted_names:
                    continue
                attempted_names.add(name)
                try:
                    sheets = list(attempt())
                    if sheets:
                        for sheet in sheets:
                            yield sheet
                        return
                except Exception as exc:
                    errors.append(f"{name} tardio: {exc}")

        if soffice:
            try:
                normalized = libreoffice_convert_to_xlsx(source, temp_root, soffice, timeout_seconds=timeout_seconds)
                for sheet in iter_openpyxl_sheets(normalized, method="libreoffice->xlsx->openpyxl"):
                    sheet.source_file = source
                    yield sheet
                return
            except Exception as exc:
                errors.append(f"LibreOffice respaldo: {exc}")

    raise ConversionError("No se pudo convertir el archivo Excel. Intentos: " + " | ".join(errors))


def build_excel_output_name(source: Path, sheet_name: str, sheet_count: int) -> str:
    base = sanitize_filename(clean_inner_tabular_extension(source.stem), fallback="archivo")
    if sheet_count <= 1:
        return f"{base}.csv"
    return f"{base}__{sanitize_filename(sheet_name, fallback='Hoja')}.csv"


def convert_excel_file(
    source: Path,
    config: RunConfig,
    outputs: OutputLocations,
    soffice: str | None,
) -> list[ConversionRecord]:
    records: list[ConversionRecord] = []

    try:
        sheets = list(iter_excel_sheets_with_fallbacks(
            source=source,
            soffice=soffice,
            timeout_seconds=config.timeout_seconds,
            prefer_libreoffice=config.prefer_libreoffice_excel,
        ))
        if not sheets:
            raise ConversionError("El archivo no contiene hojas convertibles.")
    except Exception as exc:
        return [ConversionRecord(
            kind="excel",
            source_file=str(source),
            source_ext=source.suffix.lower(),
            status="ERROR",
            error=str(exc),
        )]

    sheet_count = len(sheets)
    with tempfile.TemporaryDirectory(prefix="office_csv_write_") as tmp:
        tmp_dir = Path(tmp)
        for i, sheet in enumerate(sheets, start=1):
            item_name = sheet.sheet_name or f"Hoja{i}"
            record = ConversionRecord(
                kind="excel",
                source_file=str(source),
                source_ext=source.suffix.lower(),
                item_name=item_name,
                status="OK",
                method=sheet.method,
                warnings=" | ".join(sheet.warnings),
            )

            try:
                tmp_csv = tmp_dir / f"sheet_{i}.csv"
                rows, cols = write_rows_to_csv(
                    rows=sheet.rows,
                    output_path=tmp_csv,
                    delimiter=config.delimiter,
                    encoding=config.encoding,
                    safe_csv_mode=config.safe_csv_mode,
                )
                record.rows = rows
                record.columns = cols

                if config.output_mode in ("same", "both"):
                    same_name = build_excel_output_name(source, item_name, sheet_count)
                    target_same = source.parent / same_name
                    final_same = copy_file_unique(tmp_csv, target_same, overwrite=config.overwrite)
                    record.output_same_location = str(final_same)

                if config.output_mode in ("central", "both"):
                    if outputs.csv_dir is None:
                        raise ConversionError("No se definio carpeta central CSV.")
                    extra = item_name if sheet_count > 1 else ""
                    target_central = build_central_path(
                        outputs.csv_dir,
                        config.root_dir,
                        source,
                        suffix=".csv",
                        extra_stem=extra,
                        layout=config.central_layout,
                    )
                    final_central = copy_file_unique(tmp_csv, target_central, overwrite=config.overwrite)
                    record.output_central_location = str(final_central)

            except Exception as exc:
                record.status = "ERROR"
                record.error = str(exc)

            records.append(record)

    return records



# =========================
# Fallback universal aproximado para documentos -> PDF
# =========================


def pdf_escape_literal(data: bytes) -> bytes:
    return (
        data.replace(b"\\", b"\\\\")
        .replace(b"(", b"\\(")
        .replace(b")", b"\\)")
        .replace(b"\r", b" ")
        .replace(b"\n", b" ")
    )


def write_simple_text_pdf(lines: Sequence[str], output_path: Path, title: str = "Documento") -> Path:
    """Genera un PDF valido sin dependencias externas. Preserva texto de forma aproximada."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    page_width = 595  # A4 aprox en puntos
    page_height = 842
    margin_x = 46
    margin_y = 46
    font_size = 9
    leading = 12
    max_lines_per_page = max(1, int((page_height - 2 * margin_y) // leading))
    wrap_width = 104

    normalized_lines: list[str] = []
    header = f"PDF aproximado generado desde: {title}"
    normalized_lines.append(header)
    normalized_lines.append("-" * min(len(header), 104))
    normalized_lines.append("")

    for raw_line in lines:
        line = str(raw_line).replace("\t", "    ").rstrip()
        # HTML/XML/RTF a veces dejan espacios excesivos.
        if len(line) < 400:
            line = re.sub(r"[ \u00a0]{2,}", " ", line)
        if not line:
            normalized_lines.append("")
            continue
        wrapped = textwrap.wrap(line, width=wrap_width, break_long_words=True, replace_whitespace=False)
        normalized_lines.extend(wrapped or [""])

    if not normalized_lines:
        normalized_lines = ["Documento sin texto extraible."]

    pages: list[list[str]] = []
    for start in range(0, len(normalized_lines), max_lines_per_page):
        pages.append(normalized_lines[start:start + max_lines_per_page])
    if not pages:
        pages = [["Documento sin texto extraible."]]

    objects: dict[int, bytes] = {}
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    next_obj = 4
    page_ids: list[int] = []

    for page_lines in pages:
        content_parts: list[bytes] = []
        content_parts.append(b"BT\n")
        content_parts.append(f"/F1 {font_size} Tf\n".encode("ascii"))
        content_parts.append(f"{margin_x} {page_height - margin_y} Td\n".encode("ascii"))
        content_parts.append(f"{leading} TL\n".encode("ascii"))
        for line in page_lines:
            encoded = str(line).encode("cp1252", errors="replace")
            content_parts.append(b"(" + pdf_escape_literal(encoded) + b") Tj\nT*\n")
        content_parts.append(b"ET\n")
        content = b"".join(content_parts)

        content_id = next_obj
        objects[content_id] = b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"endstream"
        next_obj += 1

        page_id = next_obj
        page_obj = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_width} {page_height}] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("ascii")
        objects[page_id] = page_obj
        page_ids.append(page_id)
        next_obj += 1

    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("ascii")
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"

    max_obj = max(objects)
    with output_path.open("wb") as f:
        f.write(b"%PDF-1.4\n%\xE2\xE3\xCF\xD3\n")
        offsets = [0] * (max_obj + 1)
        for obj_id in range(1, max_obj + 1):
            offsets[obj_id] = f.tell()
            f.write(f"{obj_id} 0 obj\n".encode("ascii"))
            f.write(objects[obj_id])
            f.write(b"\nendobj\n")
        xref_pos = f.tell()
        f.write(f"xref\n0 {max_obj + 1}\n".encode("ascii"))
        f.write(b"0000000000 65535 f \n")
        for obj_id in range(1, max_obj + 1):
            f.write(f"{offsets[obj_id]:010d} 00000 n \n".encode("ascii"))
        f.write(f"trailer\n<< /Size {max_obj + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode("ascii"))

    validate_pdf(output_path)
    return output_path


def docx_paragraph_text(paragraph: ET.Element) -> str:
    chunks: list[str] = []
    for node in paragraph.iter():
        lname = xml_local_name(node.tag).lower()
        if lname == "t":
            chunks.append(node.text or "")
        elif lname == "tab":
            chunks.append("\t")
        elif lname in {"br", "cr"}:
            chunks.append("\n")
    text = "".join(chunks)
    text = normalize_newlines(text)
    return text


def docx_table_lines(table: ET.Element) -> list[str]:
    lines: list[str] = []
    for row in [n for n in table.iter() if xml_local_name(n.tag).lower() == "tr"]:
        cells: list[str] = []
        for cell in [n for n in list(row) if xml_local_name(n.tag).lower() == "tc"]:
            cell_parts: list[str] = []
            for p in [n for n in cell.iter() if xml_local_name(n.tag).lower() == "p"]:
                t = docx_paragraph_text(p).strip()
                if t:
                    cell_parts.append(t)
            cells.append(" / ".join(cell_parts))
        if cells:
            lines.append(" | ".join(cells))
    return lines


def extract_docx_lines(source: Path) -> tuple[list[str], list[str]]:
    warnings_out = ["PDF generado por fallback Python desde DOCX; conserva texto/tablas de forma aproximada, no el layout exacto."]
    lines: list[str] = []

    if not zipfile.is_zipfile(source):
        raise ConversionError("No es un DOCX/OOXML ZIP valido.")

    with zipfile.ZipFile(source) as zf:
        names = set(zf.namelist())
        xml_parts = []
        for candidate in sorted(names):
            lower = candidate.lower()
            if lower == "word/document.xml" or lower.startswith("word/header") or lower.startswith("word/footer"):
                if lower.endswith(".xml"):
                    xml_parts.append(candidate)
        if "word/document.xml" not in names:
            raise ConversionError("DOCX sin word/document.xml.")

        for part in xml_parts:
            try:
                root = ET.fromstring(zf.read(part))
            except Exception:
                continue
            if part.lower().startswith("word/header"):
                lines.append("[Encabezado]")
            elif part.lower().startswith("word/footer"):
                lines.append("[Pie de pagina]")

            body_candidate = xml_first_descendant(root, "body")
            body = body_candidate if body_candidate is not None else root
            for child in list(body):
                lname = xml_local_name(child.tag).lower()
                if lname == "p":
                    text = docx_paragraph_text(child)
                    for line in text.split("\n"):
                        if line.strip():
                            lines.append(line.strip())
                        elif lines and lines[-1] != "":
                            lines.append("")
                elif lname == "tbl":
                    table_out = docx_table_lines(child)
                    if table_out:
                        lines.extend(table_out)
                        lines.append("")

    if not any(line.strip() for line in lines):
        warnings_out.append("No se encontro texto visible en el DOCX; puede contener solo imagenes o estar dañado.")
    return lines, warnings_out


class SimpleHTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag_l = tag.lower()
        if tag_l in {"script", "style", "noscript"}:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag_l in {"p", "div", "br", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
            self.parts.append("\n")
        elif tag_l in {"td", "th"}:
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        tag_l = tag.lower()
        if tag_l in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag_l in {"p", "div", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)

    def get_lines(self) -> list[str]:
        text = normalize_newlines(unescape("".join(self.parts)))
        out: list[str] = []
        for line in text.split("\n"):
            cleaned = re.sub(r"[ \t\u00a0]+", " ", line).strip(" |")
            if cleaned:
                out.append(cleaned)
            elif out and out[-1] != "":
                out.append("")
        return out


def extract_html_lines(source: Path) -> tuple[list[str], list[str]]:
    text, enc = read_text_lossy(source)
    parser = SimpleHTMLTextParser()
    try:
        parser.feed(text)
        lines = parser.get_lines()
    except Exception:
        lines = [line.strip() for line in normalize_newlines(text).split("\n") if line.strip()]
    return lines, [f"PDF generado por fallback Python desde HTML/texto ({enc}); layout aproximado."]


def rtf_unicode_repl(match: re.Match[str]) -> str:
    try:
        code = int(match.group(1))
        if code < 0:
            code += 65536
        return chr(code)
    except Exception:
        return ""


def strip_rtf_to_text(text: str) -> str:
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", errors="replace"), text)
    text = re.sub(r"\\u(-?\d+)\??", rtf_unicode_repl, text)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\line", "\n", text)
    text = re.sub(r"\\tab", "\t", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("\\", "")
    return normalize_newlines(unescape(text))


def extract_rtf_lines(source: Path) -> tuple[list[str], list[str]]:
    text, enc = read_text_lossy(source)
    stripped = strip_rtf_to_text(text)
    lines = [re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in stripped.split("\n")]
    lines = [line for line in lines if line]
    return lines, [f"PDF generado por fallback Python desde RTF ({enc}); formato visual aproximado."]


def extract_odt_lines(source: Path) -> tuple[list[str], list[str]]:
    if not zipfile.is_zipfile(source):
        raise ConversionError("No es un ODT/ZIP valido.")
    lines: list[str] = []
    with zipfile.ZipFile(source) as zf:
        if "content.xml" not in zf.namelist():
            raise ConversionError("ODT sin content.xml.")
        root = ET.fromstring(zf.read("content.xml"))
        for node in root.iter():
            lname = xml_local_name(node.tag).lower()
            if lname in {"p", "h"}:
                text = re.sub(r"\s+", " ", "".join(node.itertext())).strip()
                if text:
                    lines.append(text)
            elif lname == "table-row":
                cells: list[str] = []
                for cell in [c for c in list(node) if xml_local_name(c.tag).lower() == "table-cell"]:
                    text = re.sub(r"\s+", " ", "".join(cell.itertext())).strip()
                    cells.append(text)
                if any(cells):
                    lines.append(" | ".join(cells))
    return lines, ["PDF generado por fallback Python desde ODT; conserva texto/tablas de forma aproximada."]


def run_text_extractor_command(command: str, source: Path) -> str | None:
    exe = shutil.which(command)
    if not exe:
        return None
    try:
        completed = subprocess.run(
            [exe, str(source)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=90,
        )
        text = completed.stdout or ""
        if len(text.strip()) >= 20:
            return text
    except Exception:
        return None
    return None


def printable_sequences_from_binary(data: bytes) -> list[str]:
    texts: list[str] = []
    # UTF-16LE suele encontrar texto en .doc antiguos.
    for enc in ["utf-16-le", "cp1252", "latin-1"]:
        try:
            decoded = data.decode(enc, errors="ignore")
        except Exception:
            continue
        decoded = normalize_newlines(decoded)
        decoded = re.sub(r"[^\S\r\n\t]+", " ", decoded)
        candidates = re.findall(r"[\wÁÉÍÓÚÜÑáéíóúüñ.,;:()\[\]/@#%&+\-_=¿?¡! \t]{4,}", decoded)
        for item in candidates:
            item = re.sub(r"\s+", " ", item).strip()
            if len(item) >= 4 and item not in texts:
                texts.append(item)
            if len(texts) >= 2000:
                return texts
    return texts


def extract_binary_doc_lines(source: Path) -> tuple[list[str], list[str]]:
    warnings_out = ["PDF generado por fallback Python desde documento binario; texto extraido de forma aproximada."]
    for command in ["antiword", "catdoc"]:
        text = run_text_extractor_command(command, source)
        if text:
            lines = [line.strip() for line in normalize_newlines(text).split("\n") if line.strip()]
            warnings_out.append(f"Texto extraido con {command}.")
            return lines, warnings_out

    data = source.read_bytes()
    lines = printable_sequences_from_binary(data)
    if not lines:
        lines = [
            "No se pudo extraer texto legible del documento binario.",
            "El archivo puede estar cifrado, corrupto, contener solo imagenes o requerir LibreOffice/Microsoft Word.",
        ]
    else:
        warnings_out.append("No habia motor externo; se extrajeron cadenas legibles del binario.")
    return lines, warnings_out


def extract_google_shortcut_lines(source: Path) -> tuple[list[str], list[str]]:
    warnings_out = ["El archivo es un acceso directo de Google; no contiene el documento real."]
    try:
        text, enc = read_text_lossy(source)
        data = json.loads(text)
        url = data.get("url") or data.get("doc_url") or data.get("resource_url") or ""
        doc_id = data.get("doc_id") or data.get("id") or ""
    except Exception:
        url = ""
        doc_id = ""
    lines = [
        "Este PDF es informativo.",
        f"El archivo original ({source.name}) es un acceso directo de Google Docs, no el documento real.",
        "Para conservar el contenido, abre Google Docs y descarga/exporta el documento como .docx o .pdf.",
    ]
    if url:
        lines.append(f"URL encontrada: {url}")
    if doc_id:
        lines.append(f"ID encontrado: {doc_id}")
    return lines, warnings_out


def extract_plain_text_lines(source: Path) -> tuple[list[str], list[str]]:
    text, enc = read_text_lossy(source)
    lines = [line.rstrip() for line in normalize_newlines(text).split("\n")]
    # Si parece XML generico, intenta eliminar tags de forma ligera.
    if source.suffix.lower() == ".xml" and "<" in text[:1000]:
        cleaned = re.sub(r"<[^>]+>", " ", text)
        cleaned = unescape(cleaned)
        lines = [re.sub(r"\s+", " ", line).strip() for line in normalize_newlines(cleaned).split("\n")]
    lines = [line for line in lines if line.strip()]
    return lines, [f"PDF generado por fallback Python desde texto ({enc}); layout aproximado."]


def extract_document_lines_for_python_pdf(source: Path) -> tuple[list[str], list[str]]:
    ext = source.suffix.lower()
    prefix = read_prefix(source, 8)

    if ext in GOOGLE_SHORTCUT_EXTENSIONS:
        return extract_google_shortcut_lines(source)
    if ext in {".docx", ".docm", ".dotx", ".dotm", ".docs"} or prefix.startswith(ZIP_SIGNATURES):
        try:
            return extract_docx_lines(source)
        except Exception:
            if zipfile.is_zipfile(source):
                try:
                    return extract_odt_lines(source)
                except Exception:
                    pass
            # Puede ser texto con extension rara.
            if is_probably_text_bytes(read_prefix(source, 131072)):
                return extract_plain_text_lines(source)
            raise
    if ext in {".odt", ".ott", ".fodt", ".sxw"}:
        try:
            return extract_odt_lines(source)
        except Exception:
            if is_probably_text_bytes(read_prefix(source, 131072)):
                return extract_plain_text_lines(source)
            raise
    if ext in {".html", ".htm", ".xhtml"}:
        return extract_html_lines(source)
    if ext == ".rtf":
        return extract_rtf_lines(source)
    if ext in {".txt", ".text", ".xml", ".abw", ".wri"} or is_probably_text_bytes(read_prefix(source, 131072)):
        return extract_plain_text_lines(source)
    return extract_binary_doc_lines(source)


def python_text_document_to_pdf(source: Path, temp_root: Path) -> tuple[Path, list[str]]:
    try:
        lines, warnings_out = extract_document_lines_for_python_pdf(source)
    except Exception as exc:
        lines = [
            "No se pudo extraer el contenido del documento con el fallback Python.",
            f"Archivo: {source}",
            f"Error: {exc}",
            "Prueba instalando LibreOffice para una conversion fiel a PDF.",
        ]
        warnings_out = [f"Fallback Python genero un PDF informativo porque no pudo extraer texto: {exc}"]

    if not any(str(line).strip() for line in lines):
        lines = [
            "Documento sin texto visible extraible por el fallback Python.",
            "Puede contener solo imagenes, estar cifrado, corrupto o requerir LibreOffice/Microsoft Word.",
        ]
        warnings_out.append("No se encontro texto visible; se genero PDF informativo.")

    outdir = temp_root / "python_text_pdf"
    out_pdf = outdir / f"{sanitize_filename(source.stem)}.pdf"
    write_simple_text_pdf(lines, out_pdf, title=source.name)
    return out_pdf, warnings_out


# =========================
# Documentos -> PDF
# =========================

def is_google_shortcut(source: Path) -> bool:
    return source.suffix.lower() in GOOGLE_SHORTCUT_EXTENSIONS


def produce_document_pdf(
    source: Path,
    config: RunConfig,
    soffice: str | None,
    temp_parent: Path,
) -> tuple[Path, str, list[str]]:
    warnings: list[str] = []
    errors: list[str] = []

    if is_google_shortcut(source):
        warnings.append(
            f"{source.suffix} es un acceso directo de Google; el fallback generara un PDF informativo "
            "porque el archivo no contiene el documento real."
        )

    # .pages normalmente no es soportado por LibreOffice en muchos entornos.
    if source.suffix.lower() == ".pages":
        warnings.append("Formato .pages: normalmente debe exportarse desde Apple Pages a .docx o .pdf antes de convertir.")

    method_order: list[str] = []
    if config.prefer_ms_word and microsoft_word_com_available():
        method_order.append("ms_word")
    if soffice and not is_google_shortcut(source):
        method_order.append("libreoffice")
    if "ms_word" not in method_order and microsoft_word_com_available() and not is_google_shortcut(source):
        method_order.append("ms_word")

    # Ultimo recurso: PDF aproximado generado 100% en Python.
    # Este fallback evita que un .docx falle solo porque no existe LibreOffice en Android/Termux.
    method_order.append("python_text_pdf")

    for method in method_order:
        engine_tmp = temp_parent / f"engine_{method}_{hashlib.sha1(str(source).encode()).hexdigest()[:8]}"
        engine_tmp.mkdir(parents=True, exist_ok=True)
        try:
            method_warnings: list[str] = []
            if method == "libreoffice":
                if not soffice:
                    raise ConversionError("LibreOffice no detectado.")
                pdf = libreoffice_convert_document_to_pdf(
                    source=source,
                    temp_root=engine_tmp,
                    soffice=soffice,
                    timeout_seconds=config.timeout_seconds,
                    pdfa=config.pdfa,
                )
                method_name = "libreoffice->pdf"
            elif method == "ms_word":
                pdf = ms_word_convert_to_pdf(source=source, temp_root=engine_tmp, pdfa=config.pdfa)
                method_name = "microsoft-word-com->pdf"
            elif method == "python_text_pdf":
                pdf, method_warnings = python_text_document_to_pdf(source=source, temp_root=engine_tmp)
                method_name = "python-text-fallback->pdf"
            else:
                raise ConversionError(f"Motor desconocido: {method}")

            validate_pdf(pdf)
            persistent_pdf = temp_parent / f"result_{hashlib.sha1((str(source)+method).encode()).hexdigest()[:12]}.pdf"
            shutil.copy2(pdf, persistent_pdf)
            return persistent_pdf, method_name, warnings + method_warnings

        except Exception as exc:
            errors.append(f"{method}: {exc}")

    # Si incluso el fallback falla, genera un PDF de emergencia con el detalle de errores.
    try:
        emergency_lines = [
            "No se pudo convertir fielmente este documento.",
            f"Archivo: {source}",
            "Intentos realizados:",
            *errors,
        ]
        emergency_pdf = temp_parent / "emergency_error.pdf"
        write_simple_text_pdf(emergency_lines, emergency_pdf, title=source.name)
        return emergency_pdf, "python-emergency-error-pdf", warnings + ["Se genero PDF informativo de emergencia con los errores."]
    except Exception:
        raise ConversionError("No se pudo convertir el documento a PDF. Intentos: " + " | ".join(errors))

def build_document_same_output_path(source: Path) -> Path:
    return source.parent / f"{sanitize_filename(source.stem)}.pdf"


def convert_document_file(
    source: Path,
    config: RunConfig,
    outputs: OutputLocations,
    soffice: str | None,
) -> list[ConversionRecord]:
    record = ConversionRecord(
        kind="docs",
        source_file=str(source),
        source_ext=source.suffix.lower(),
        item_name=source.name,
        status="OK",
    )

    with tempfile.TemporaryDirectory(prefix="office_doc_pdf_final_") as tmp:
        temp_parent = Path(tmp)
        try:
            tmp_pdf, method, warnings = produce_document_pdf(source, config, soffice, temp_parent=temp_parent)
            record.method = method
            record.warnings = " | ".join(warnings)

            if config.output_mode in ("same", "both"):
                target_same = build_document_same_output_path(source)
                final_same = copy_file_unique(tmp_pdf, target_same, overwrite=config.overwrite)
                record.output_same_location = str(final_same)

            if config.output_mode in ("central", "both"):
                if outputs.pdf_dir is None:
                    raise ConversionError("No se definio carpeta central PDF.")
                target_central = build_central_path(
                    outputs.pdf_dir,
                    config.root_dir,
                    source,
                    suffix=".pdf",
                    extra_stem="",
                    layout=config.central_layout,
                )
                final_central = copy_file_unique(tmp_pdf, target_central, overwrite=config.overwrite)
                record.output_central_location = str(final_central)

        except Exception as exc:
            record.status = "ERROR"
            record.error = str(exc)

    return [record]


# =========================
# Descubrimiento y reportes
# =========================

def discover_jobs(root_dir: Path, convert_choice: ConvertChoice, central_root: Path | None = None) -> list[JobFile]:
    include_excel = convert_choice in ("excel", "both")
    include_docs = convert_choice in ("docs", "both")
    jobs: list[JobFile] = []

    for path in root_dir.rglob("*"):
        if not path.is_file():
            continue
        if should_skip_path(path, central_root=central_root):
            continue
        ext = path.suffix.lower()

        # Si una extension esta en ambos, damos prioridad a Excel solo cuando corresponde claramente.
        # .xml es ambiguo; al elegir ambos se intentara como Excel primero por ser frecuente en Excel 2003 XML.
        if include_excel and ext in EXCEL_EXTENSIONS:
            jobs.append(JobFile(kind="excel", path=path))
            continue
        if include_docs and ext in DOCUMENT_EXTENSIONS:
            jobs.append(JobFile(kind="docs", path=path))
            continue

    return sorted(jobs, key=lambda j: (j.kind, str(j.path).lower()))


def create_output_locations(config: RunConfig, stamp: str) -> OutputLocations:
    if config.output_mode not in ("central", "both"):
        reports_dir = Path.cwd().resolve()
        return OutputLocations(reports_dir=reports_dir)

    central_root = Path.cwd().resolve() / f"{CENTRAL_FOLDER_PREFIX}_{stamp}"
    reports_dir = central_root / "REPORTES"
    failed_dir = central_root / "NO_CONVERTIDOS"
    csv_dir = central_root / "CSV"
    pdf_dir = central_root / "PDF"

    central_root.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)
    failed_dir.mkdir(parents=True, exist_ok=True)

    if config.convert_choice in ("excel", "both"):
        csv_dir.mkdir(parents=True, exist_ok=True)
    else:
        csv_dir = None

    if config.convert_choice in ("docs", "both"):
        pdf_dir.mkdir(parents=True, exist_ok=True)
    else:
        pdf_dir = None

    return OutputLocations(
        central_root=central_root,
        csv_dir=csv_dir,
        pdf_dir=pdf_dir,
        reports_dir=reports_dir,
        failed_dir=failed_dir,
    )


def copy_failed_original_if_needed(
    record: ConversionRecord,
    source: Path,
    config: RunConfig,
    outputs: OutputLocations,
) -> None:
    if record.status == "OK":
        return
    if not config.copy_failed_originals:
        return
    if outputs.failed_dir is None:
        return

    try:
        kind_dir = outputs.failed_dir / ("EXCEL" if record.kind == "excel" else "DOCUMENTOS")
        target = build_central_path(
            kind_dir,
            config.root_dir,
            source,
            suffix=source.suffix,
            extra_stem="",
            layout=config.central_layout,
        )
        copied = copy_file_unique(source, target, overwrite=False)
        record.failed_original_copy = str(copied)
    except Exception as exc:
        record.warnings = (record.warnings + " | " if record.warnings else "") + f"No se pudo copiar original fallido: {exc}"


def write_report(records: list[ConversionRecord], report_path: Path, encoding: str = DEFAULT_ENCODING) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "kind",
        "source_file",
        "source_ext",
        "item_name",
        "status",
        "method",
        "rows",
        "columns",
        "output_same_location",
        "output_central_location",
        "failed_original_copy",
        "warnings",
        "error",
    ]
    with report_path.open("w", newline="", encoding=encoding) as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for record in records:
            writer.writerow(record.__dict__)


def write_error_log(records: list[ConversionRecord], error_log_path: Path) -> None:
    errors = [r for r in records if r.status != "OK"]
    if not errors:
        return
    error_log_path.parent.mkdir(parents=True, exist_ok=True)
    with error_log_path.open("w", encoding="utf-8") as f:
        for record in errors:
            f.write(f"Tipo: {record.kind}\n")
            f.write(f"Archivo: {record.source_file}\n")
            if record.item_name:
                f.write(f"Elemento: {record.item_name}\n")
            f.write(f"Error: {record.error}\n")
            if record.failed_original_copy:
                f.write(f"Copia del original fallido: {record.failed_original_copy}\n")
            f.write("-" * 90 + "\n")


def write_summary_txt(records: list[ConversionRecord], summary_path: Path, outputs: OutputLocations, config: RunConfig) -> None:
    total_ok = sum(1 for r in records if r.status == "OK")
    total_errors = sum(1 for r in records if r.status != "OK")
    excel_ok = sum(1 for r in records if r.kind == "excel" and r.status == "OK")
    docs_ok = sum(1 for r in records if r.kind == "docs" and r.status == "OK")

    lines = [
        "Resumen de conversion",
        "=====================",
        f"Fecha: {_dt.datetime.now().isoformat(sep=' ', timespec='seconds')}",
        f"Carpeta raiz recorrida: {config.root_dir}",
        f"Modo de conversion: {config.convert_choice}",
        f"Modo de salida: {config.output_mode}",
        f"Resultados OK: {total_ok}",
        f"Errores: {total_errors}",
        f"Hojas Excel convertidas a CSV: {excel_ok}",
        f"Documentos convertidos a PDF: {docs_ok}",
    ]
    if outputs.central_root:
        lines.append(f"Carpeta central: {outputs.central_root}")
    if outputs.csv_dir:
        lines.append(f"Subcarpeta CSV: {outputs.csv_dir}")
    if outputs.pdf_dir:
        lines.append(f"Subcarpeta PDF: {outputs.pdf_dir}")
    if outputs.failed_dir:
        lines.append(f"Originales fallidos: {outputs.failed_dir}")

    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# =========================
# Interfaz CLI / preguntas
# =========================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convierte recursivamente Excel a CSV y documentos Word/Writer a PDF.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("root", nargs="?", default=None, help="Carpeta principal a recorrer. Si se omite, se pregunta o se usa la actual.")
    parser.add_argument("--convert", choices=["excel", "docs", "both"], default=None, help="Que convertir: excel, docs o both.")
    parser.add_argument("--output", choices=["same", "central", "both"], default=None, help="Donde guardar: same, central o both.")
    parser.add_argument("--central-layout", choices=["flat", "tree"], default="flat", help="Organizacion dentro de CSV/PDF central: flat o tree.")
    parser.add_argument("--delimiter", default=DEFAULT_DELIMITER, help="Delimitador del CSV, por ejemplo ',' o ';'.")
    parser.add_argument("--encoding", default=DEFAULT_ENCODING, help="Codificacion de CSV y reportes.")
    parser.add_argument("--overwrite", action="store_true", help="Sobrescribe salidas existentes. Si no, crea nombres _2, _3, etc.")
    parser.add_argument("--safe-csv", action="store_true", help="Escapa textos que parecen formulas al abrir CSV en Excel.")
    parser.add_argument("--prefer-libreoffice-excel", action="store_true", help="Normaliza/recalcula Excel con LibreOffice antes de leer con Python.")
    parser.add_argument("--no-libreoffice", action="store_true", help="No usa LibreOffice aunque este instalado; util si se cuelga o quieres usar fallbacks Python.")
    parser.add_argument("--prefer-ms-word", action="store_true", help="En Windows, intenta Microsoft Word COM antes que LibreOffice para documentos.")
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SECONDS, help="Tiempo maximo por conversion con LibreOffice, en segundos.")
    parser.add_argument("--pdfa", choices=["none", "1", "2", "3"], default="none", help="Exportar PDF normal o PDF/A-1/2/3 cuando sea posible.")
    parser.add_argument("--no-copy-failed", action="store_true", help="No copia originales fallidos a NO_CONVERTIDOS.")
    parser.add_argument("--non-interactive", action="store_true", help="No pregunta nada; usa argumentos o valores por defecto.")
    return parser.parse_args()


def ask_root(default: Path) -> Path:
    raw = input(f"Carpeta principal a recorrer [Enter = {default}]: ").strip().strip('"')
    return Path(raw).expanduser().resolve() if raw else default


def ask_convert_choice() -> ConvertChoice:
    print("\n¿Que quieres convertir?")
    print("  1) Solo Excel/Calc a CSV")
    print("  2) Solo documentos Word/Writer a PDF")
    print("  3) Ambos: Excel a CSV y documentos a PDF")
    while True:
        raw = input("Elige 1, 2 o 3: ").strip()
        if raw == "1":
            return "excel"
        if raw == "2":
            return "docs"
        if raw == "3":
            return "both"
        print("Opcion invalida. Escribe 1, 2 o 3.")


def ask_output_mode() -> OutputMode:
    print("\n¿Donde quieres guardar los archivos convertidos?")
    print("  1) En la misma carpeta donde esta cada original")
    print("  2) En una carpeta central creada donde ejecutes el script")
    print("  3) Ambas opciones")
    while True:
        raw = input("Elige 1, 2 o 3: ").strip()
        if raw == "1":
            return "same"
        if raw == "2":
            return "central"
        if raw == "3":
            return "both"
        print("Opcion invalida. Escribe 1, 2 o 3.")


def build_config_from_args(args: argparse.Namespace) -> RunConfig:
    cwd = Path.cwd().resolve()

    if args.non_interactive:
        root_dir = Path(args.root).expanduser().resolve() if args.root else cwd
        convert_choice: ConvertChoice = args.convert or "both"
        output_mode: OutputMode = args.output or "both"
    else:
        root_dir = Path(args.root).expanduser().resolve() if args.root else ask_root(cwd)
        convert_choice = args.convert or ask_convert_choice()
        output_mode = args.output or ask_output_mode()

    if not root_dir.exists() or not root_dir.is_dir():
        raise ConversionError(f"La carpeta no existe o no es valida: {root_dir}")

    delimiter = args.delimiter
    if len(delimiter) != 1:
        raise ConversionError("El delimitador CSV debe ser un solo caracter, por ejemplo ',' o ';'.")

    return RunConfig(
        root_dir=root_dir,
        convert_choice=convert_choice,
        output_mode=output_mode,
        delimiter=delimiter,
        encoding=args.encoding,
        overwrite=bool(args.overwrite),
        safe_csv_mode=bool(args.safe_csv),
        prefer_libreoffice_excel=bool(args.prefer_libreoffice_excel),
        disable_libreoffice=bool(args.no_libreoffice),
        prefer_ms_word=bool(args.prefer_ms_word),
        timeout_seconds=int(args.timeout),
        central_layout=args.central_layout,
        pdfa=args.pdfa,
        copy_failed_originals=not bool(args.no_copy_failed),
        non_interactive=bool(args.non_interactive),
    )


# =========================
# Programa principal
# =========================

def main() -> int:
    try:
        sys.stdout.reconfigure(line_buffering=True)
        sys.stderr.reconfigure(line_buffering=True)
    except Exception:
        pass

    args = parse_args()
    stamp = now_stamp()

    try:
        config = build_config_from_args(args)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    soffice = None if config.disable_libreoffice else get_soffice_executable()

    print_header("Convertidor Office -> CSV/PDF")
    print(f"Carpeta raiz: {config.root_dir}")
    print(f"Conversion: {config.convert_choice}")
    print(f"Salida: {config.output_mode}")

    if config.disable_libreoffice:
        print("LibreOffice desactivado por --no-libreoffice.")
    elif soffice:
        print(f"LibreOffice detectado: {soffice}")
    else:
        print("Aviso: LibreOffice no fue detectado.")
        print("       Excel se intentara con librerias Python y parsers de archivos disfrazados.")
        print("       Documentos a PDF usaran fallback Python aproximado; instala LibreOffice para PDF fiel.")

    if microsoft_word_com_available():
        print("Microsoft Word COM detectado en Windows como respaldo para PDF.")

    jobs = discover_jobs(config.root_dir, config.convert_choice, central_root=None)

    if not jobs:
        print("\nNo se encontraron archivos compatibles para la opcion elegida.")
        return 0

    outputs = create_output_locations(config, stamp)
    if outputs.central_root:
        print(f"Carpeta central: {outputs.central_root}")
        if outputs.csv_dir:
            print(f"  CSV: {outputs.csv_dir}")
        if outputs.pdf_dir:
            print(f"  PDF: {outputs.pdf_dir}")

    print(f"\nArchivos encontrados: {len(jobs)}")
    all_records: list[ConversionRecord] = []

    for idx, job in enumerate(jobs, start=1):
        source = job.path
        print(f"[{idx}/{len(jobs)}] {job.kind.upper()} -> {source}")

        try:
            if job.kind == "excel":
                records = convert_excel_file(source, config, outputs, soffice=soffice)
            elif job.kind == "docs":
                records = convert_document_file(source, config, outputs, soffice=soffice)
            else:
                records = [ConversionRecord(kind=job.kind, source_file=str(source), source_ext=source.suffix.lower(), status="ERROR", error="Tipo de trabajo desconocido.")]

            for record in records:
                copy_failed_original_if_needed(record, source, config, outputs)
            all_records.extend(records)

            ok_count = sum(1 for r in records if r.status == "OK")
            error_count = sum(1 for r in records if r.status != "OK")
            print(f"    OK: {ok_count} | Errores: {error_count}")

        except KeyboardInterrupt:
            print("\nProceso cancelado por el usuario.")
            break
        except Exception as exc:
            record = ConversionRecord(
                kind=job.kind,
                source_file=str(source),
                source_ext=source.suffix.lower(),
                status="ERROR",
                error=f"Error inesperado: {exc}\n{traceback.format_exc()}",
            )
            copy_failed_original_if_needed(record, source, config, outputs)
            all_records.append(record)
            print(f"    ERROR inesperado: {exc}")

    reports_dir = outputs.reports_dir or Path.cwd().resolve()
    report_path = reports_dir / f"{REPORT_BASENAME}_{stamp}.csv"
    error_log_path = reports_dir / f"{ERROR_LOG_BASENAME}_{stamp}.log"
    summary_path = reports_dir / f"resumen_conversion_{stamp}.txt"

    write_report(all_records, report_path, encoding=config.encoding)
    write_error_log(all_records, error_log_path)
    write_summary_txt(all_records, summary_path, outputs, config)

    total_ok = sum(1 for r in all_records if r.status == "OK")
    total_errors = sum(1 for r in all_records if r.status != "OK")
    excel_ok = sum(1 for r in all_records if r.kind == "excel" and r.status == "OK")
    docs_ok = sum(1 for r in all_records if r.kind == "docs" and r.status == "OK")

    print_header("Resumen final")
    print(f"Resultados OK: {total_ok}")
    print(f"  Hojas Excel convertidas a CSV: {excel_ok}")
    print(f"  Documentos convertidos a PDF: {docs_ok}")
    print(f"Errores: {total_errors}")
    print(f"Reporte CSV: {report_path}")
    print(f"Resumen TXT: {summary_path}")
    if total_errors:
        print(f"Log de errores: {error_log_path}")
        if outputs.failed_dir:
            print(f"Originales fallidos copiados en: {outputs.failed_dir}")
    if outputs.central_root:
        print(f"Carpeta central: {outputs.central_root}")

    return 0 if total_errors == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
