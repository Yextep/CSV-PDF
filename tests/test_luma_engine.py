import csv
import importlib.util
import json
import sys
import zipfile
from pathlib import Path
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("luma_engine", ROOT / "android/app/src/main/python/luma_engine.py")
engine = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = engine
spec.loader.exec_module(engine)


def run(source, out, **options):
    return json.loads(engine.convert(str(source), str(out), json.dumps(options)))


def test_xlsx_sheets_unicode_formula_and_delimiter(tmp_path):
    source = tmp_path / "datos.xlsx"
    book = Workbook()
    book.active.title = "Café"
    book.active.append(["Nombre", "Valor", "Texto"])
    book.active.append(["Muñoz", 42, "+123"])
    book.active.append(["Fórmula", "=SUM(1,2)", "a;b"])
    book.create_sheet("Resumen").append([True, 1.5])
    book.save(source)
    result = run(source, tmp_path / "out", delimiter=";", safe=True)
    assert len(result["files"]) == 2
    with (tmp_path / "out" / result["files"][0]).open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file, delimiter=";"))
    assert rows[1] == ["Muñoz", "42", "'+123"]
    assert rows[2][1] == "'=SUM(1,2)"
    assert rows[2][2] == "a;b"
    assert any("formula" in w.lower() for w in result["warnings"])


def test_csv_collision_and_source_is_untouched(tmp_path):
    engine.create_demo(str(tmp_path))
    source = tmp_path / "Pequeño inventario.xlsx"
    before = source.read_bytes()
    first = run(source, tmp_path / "out")
    second = run(source, tmp_path / "out")
    assert first["files"] != second["files"]
    assert source.read_bytes() == before


def test_docx_text_and_table(tmp_path):
    source = tmp_path / "documento.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Hola, niña — café</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Una celda</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
    result = run(source, tmp_path / "out")
    assert result["kind"] == "pdf"
    assert "Hola, niña — café" in result["lines"]
    assert "Una celda" in result["lines"]
    assert result["warnings"]


def test_empty_and_corrupt_documents_do_not_generate_emergency_pdf(tmp_path):
    for name, data in [("roto.docx", b"not a zip"), ("vacio.txt", b""), ("blanco.txt", b" \n\n"), ("antiguo.doc", b"binary")]:
        source = tmp_path / name
        source.write_bytes(data)
        with pytest.raises(Exception):
            run(source, tmp_path / "out")
    assert not list(tmp_path.rglob("*.pdf"))


def test_utf16_and_rtf_unicode(tmp_path):
    source = tmp_path / "texto.txt"
    source.write_text("España, café y té", encoding="utf-16")
    assert run(source, tmp_path / "out")["lines"] == ["España, café y té"]
    source = tmp_path / "texto.rtf"
    source.write_text(r"{\rtf1\ansi Caf\'e9 y ni\u241?a\par Segunda l\'ednea}")
    assert "Café y niña" in run(source, tmp_path / "rtf")["lines"]


ODF = '<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet><table:table table:name="Hoja"><table:table-row><table:table-cell office:value-type="string"><text:p>Café</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="3.5"/><table:table-cell table:number-columns-repeated="1000000"/></table:table-row><table:table-row table:number-rows-repeated="1000000"><table:table-cell table:number-columns-repeated="1000000"/></table:table-row></table:table></office:spreadsheet></office:body></office:document>'


@pytest.mark.parametrize("extension", ["ods", "fods"])
def test_odf_ignores_million_empty_repeated_cells(tmp_path, extension):
    source = tmp_path / ("datos." + extension)
    if extension == "ods":
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("content.xml", ODF)
    else:
        source.write_text(ODF)
    result = run(source, tmp_path / "out")
    assert (tmp_path / "out" / result["files"][0]).read_text(encoding="utf-8-sig") == "Café,3.5\n"


def test_transaction_removes_partial_csv_on_encoding_failure(tmp_path):
    book = Workbook()
    book.active.append(["漢字"])
    source = tmp_path / "unicode.xlsx"
    book.save(source)
    with pytest.raises(UnicodeEncodeError):
        run(source, tmp_path / "out", encoding="cp1252")
    assert not list((tmp_path / "out").iterdir())


def test_excessive_spreadsheet_dimensions_and_xml_entities(tmp_path):
    book = Workbook()
    book.active.cell(1000000, 1, "lejos")
    source = tmp_path / "enorme.xlsx"
    book.save(source)
    with pytest.raises(ValueError, match="500.000"):
        run(source, tmp_path / "out")
    source = tmp_path / "entidad.fods"
    source.write_text('<!DOCTYPE x [<!ENTITY a "peligro">]><x>&a;</x>')
    with pytest.raises(Exception):
        run(source, tmp_path / "out")


def test_zip_size_limit_and_options(tmp_path, monkeypatch):
    source = tmp_path / "grande.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/document.xml", "text")
    original = zipfile.ZipFile.infolist
    def huge(archive):
        info = original(archive)
        info[0].file_size = 101 * 1024 * 1024
        return info
    monkeypatch.setattr(zipfile.ZipFile, "infolist", huge)
    with pytest.raises(ValueError, match="expandido"):
        run(source, tmp_path / "out")
    with pytest.raises(ValueError, match="configuración"):
        run(source, tmp_path / "out", delimiter="dos")
